from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Assignment4_Report_Amirkhan_Tulkibayev.docx"
UML_IMAGE = ROOT / "docs" / "uml-class-diagram.png"
GITHUB_URL = "https://github.com/hann-qaz/Assignment4_facade_pattern"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="8"):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    row_properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_properties.append(header)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    run = paragraph.add_run()
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def add_code(document, code):
    paragraph = document.add_paragraph(style="Code")
    lines = code.strip("\n").splitlines()
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        if line.strip().startswith("//"):
            run.font.color.rgb = RGBColor(31, 78, 121)
            run.font.italic = True
        if index < len(lines) - 1:
            run.add_break()
    return paragraph


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)

    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E79")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F4F8FC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run(value)
    return table


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def build_report():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    title_style = document.styles["Title"]
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(18)
    title_properties = title_style._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    code_style = document.styles.add_style("Code", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor(31, 41, 55)
    code_style.paragraph_format.left_indent = Inches(0.25)
    code_style.paragraph_format.right_indent = Inches(0.15)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(7)
    code_style.paragraph_format.line_spacing = 1.0
    code_style.paragraph_format.keep_together = True

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run("Amirkhan Tulkibayev  |  SE-2522  |  Page ")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(89, 89, 89)
    add_page_number(footer_paragraph)

    # Cover page
    document.add_paragraph("Assignment 4 Facade Pattern Report", style="Title")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle_run = subtitle.add_run("Online Order Processing Facade")
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(15)
    subtitle.paragraph_format.space_after = Pt(28)

    details = [
        ("Student", "Amirkhan Tulkibayev"),
        ("Group", "SE-2522"),
        ("Course", "Software Design Patterns"),
        ("Language", "Java 17"),
        ("Pattern", "Facade"),
    ]
    details_table = add_table(
        document,
        ["Field", "Value"],
        details,
        [Inches(1.55), Inches(5.25)],
    )
    document.add_paragraph()
    link_paragraph = document.add_paragraph()
    link_paragraph.add_run("GitHub repository  ").bold = True
    add_hyperlink(link_paragraph, GITHUB_URL, GITHUB_URL)

    summary = document.add_paragraph()
    summary.paragraph_format.space_before = Pt(18)
    summary.add_run("Project summary  ").bold = True
    summary.add_run(
        "The project provides a small API for placing and cancelling online orders. "
        "The Facade coordinates five independent subsystem classes and performs rollback when payment or shipping fails."
    )

    document.add_page_break()

    # Introduction
    document.add_heading("1 Introduction", level=1)
    document.add_paragraph(
        "I chose an online order processing subsystem. Placing one order requires several dependent steps: "
        "reserve stock, charge the customer, schedule delivery, save the order, and send a notification. "
        "A direct client must know this order and undo completed steps after a failure. OrderFacade instead exposes placeOrder and cancelOrder. "
        "It hides internal identifiers and rollback order, while the public subsystem classes remain available to advanced clients."
    )

    document.add_heading("Client comparison", level=2)
    before = document.add_paragraph()
    before.add_run("Without a Facade  ").bold = True
    before.add_run("the client would coordinate every step and every compensation action.")
    add_code(
        document,
        "inventory.reserve(orderId, items);\n"
        "paymentId = payment.charge(orderId, token, total);\n"
        "shipmentId = shipping.scheduleDelivery(orderId, address);",
    )
    after = document.add_paragraph()
    after.add_run("With the Facade  ").bold = True
    after.add_run("the client states the business action in one call.")
    add_code(document, "OrderResult result = orderFacade.placeOrder(cart, customer);")

    document.add_heading("Responsibilities", level=2)
    add_table(
        document,
        ["Class", "Responsibility"],
        [
            ("OrderFacade", "Coordinates the workflow and rollback; exposes the small client API."),
            ("InventoryService", "Checks stock, creates reservations, and releases reserved items."),
            ("PaymentProcessor", "Charges payments and records refunds."),
            ("ShippingService", "Schedules and cancels deliveries."),
            ("NotificationService", "Sends confirmation, cancellation, and failure messages."),
            ("OrderRepository", "Stores confirmed orders and updates their status."),
        ],
        [Inches(2.0), Inches(4.8)],
    )

    document.add_page_break()

    # UML
    document.add_heading("2 UML Class Diagram", level=1)
    document.add_paragraph(
        "DemoClient depends on OrderFacade. The Facade holds references to the subsystem services and delegates work to them. "
        "The one-way relationships are important: subsystem classes do not store a reference to the Facade and do not know that it exists."
    )
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.keep_together = True
    uml_shape = picture_paragraph.add_run().add_picture(str(UML_IMAGE), width=Inches(6.3))
    uml_shape._inline.docPr.set(
        "descr",
        "UML class diagram showing DemoClient using OrderFacade, which coordinates five independent subsystem classes.",
    )
    caption = document.add_paragraph("Figure 1  Facade pattern class diagram")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(9)

    document.add_heading("Workflow and rollback", level=2)
    add_table(
        document,
        ["Situation", "Facade action"],
        [
            ("Success", "Reserve stock, charge payment, schedule shipping, save order, notify customer."),
            ("Insufficient stock", "Stop before payment and notify the customer."),
            ("Payment declined", "Release the stock reservation and notify the customer."),
            ("Shipping failed", "Refund payment, release stock, and notify the customer."),
            ("Cancellation", "Cancel shipping, refund payment, release stock, update status, and notify."),
        ],
        [Inches(1.7), Inches(5.1)],
    )

    document.add_page_break()

    # Clean Code
    document.add_heading("3 Clean Code Principles", level=1)
    document.add_paragraph(
        "The following principles are visible in the implementation. Each excerpt includes a short annotation that explains the design choice."
    )

    document.add_heading("3.1 Meaningful names", level=2)
    document.add_paragraph(
        "Names describe business actions instead of technical steps. A reader can understand the client code without opening subsystem classes."
    )
    add_code(
        document,
        "OrderResult placedOrder = orderFacade.placeOrder(cart, customer);\n"
        "// placeOrder states the full business intention.",
    )

    document.add_heading("3.2 Single Responsibility", level=2)
    document.add_paragraph(
        "OrderFacade coordinates the workflow but does not calculate stock, process card data, or create delivery records. It delegates each task."
    )
    add_code(
        document,
        "inventoryService.reserve(orderId, cart.items());\n"
        "paymentId = paymentProcessor.charge(orderId, customer.paymentToken(), cart.total());\n"
        "// The Facade coordinates; each service performs its own work.",
    )

    document.add_heading("3.3 No leaking of subsystem details", level=2)
    document.add_paragraph(
        "The public method accepts simple domain objects and returns OrderResult. It does not return PaymentRecord, Shipment, or a repository object."
    )
    add_code(
        document,
        "public OrderResult placeOrder(Cart cart, Customer customer)\n"
        "// The client receives one stable result type, not subsystem types.",
    )

    document.add_heading("3.4 Small focused subsystem classes", level=2)
    document.add_paragraph(
        "Each subsystem class has one clear reason to change. For example, payment rules can change without changing inventory or notification code."
    )
    add_code(
        document,
        "public String charge(String orderId, String paymentToken, BigDecimal amount)\n"
        "public boolean refund(String paymentId)\n"
        "// PaymentProcessor contains payment operations only.",
    )

    document.add_heading("3.5 Explicit error handling and rollback", level=2)
    document.add_paragraph(
        "Known subsystem failures are handled where coordination happens. Completed steps are compensated in reverse order so the system does not keep money or stock after a failed delivery."
    )
    add_code(
        document,
        "catch (ShippingException exception) {\n"
        "    paymentProcessor.refund(paymentId);\n"
        "    inventoryService.release(orderId);\n"
        "    return failOrder(customer, orderId, exception.getMessage());\n"
        "}\n"
        "// Refund and release restore a consistent state.",
    )

    document.add_page_break()

    # Tests and conclusion
    document.add_heading("4 Testing", level=1)
    document.add_paragraph(
        "Seven JUnit 5 tests verify the normal workflow, the rollback paths, and important edge cases. "
        "Each test creates fresh in-memory services, so the tests are fast, independent, repeatable, and self-validating."
    )
    add_table(
        document,
        ["Test scenario", "Expected result"],
        [
            ("Successful order", "All subsystem results and the confirmation are correct."),
            ("Payment declined", "Stock is restored; shipping is not scheduled."),
            ("Shipping failure", "Payment and stock are restored."),
            ("Insufficient stock", "Processing stops before payment."),
            ("Cancellation", "Resources are restored and status becomes cancelled."),
            ("Second cancellation", "The request is rejected without repeated rollback."),
            ("Repeated product lines", "Combined quantity is checked before reservation."),
        ],
        [Inches(2.0), Inches(4.8)],
    )

    document.add_heading("5 Conclusion", level=1)
    document.add_paragraph(
        "OrderFacade simplifies a real multi-step subsystem into two client methods. It hides construction order, service call order, internal identifiers, "
        "persistence steps, notification selection, and rollback. The Facade does not replace or block the subsystem. InventoryService, PaymentProcessor, "
        "ShippingService, NotificationService, and OrderRepository remain independent and directly accessible when an advanced client needs lower-level control."
    )
    document.add_paragraph(
        "This design follows the lecture definition of Facade: a unified, higher-level interface makes a set of subsystem interfaces easier to use while reducing client coupling."
    )

    document.add_heading("Repository", level=2)
    repository_link = document.add_paragraph()
    add_hyperlink(repository_link, GITHUB_URL, GITHUB_URL)

    document.add_heading("References", level=2)
    references = [
        "Assignment 4 Facade Pattern Instructions",
        "Lecture 6  The Facade Design Pattern",
    ]
    for item in references:
        document.add_paragraph(item, style="List Bullet")

    document.core_properties.title = "Assignment 4 Facade Pattern Report"
    document.core_properties.subject = "Online Order Processing Facade in Java"
    document.core_properties.author = "Amirkhan Tulkibayev"
    document.core_properties.keywords = "Facade, Java, Clean Code, JUnit, UML"

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
